"""RAG behavioral coverage with real database, parsers, ranking and vector storage."""
from io import BytesIO
from unittest.mock import patch
from zipfile import ZipFile

import pytest
from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from rag.models import Chunk, Document, DocumentCollection, RAGPipeline, VectorStore
from rag.parsing import extract_file_content, fetch_public_text
from rag.pipeline import DocumentProcessor, EmbeddingService, RAGService, RAGUnavailable
from rag.retrieval import split_text

pytestmark = pytest.mark.django_db
BASE = '/api/v1/rag'


@pytest.fixture
def collection(user):
    store = VectorStore.objects.create(name='Local', store_type='memory', owner=user)
    return DocumentCollection.objects.create(name='运维知识', vector_store=store, owner=user, chunk_size=180, chunk_overlap=25)


def add(collection, text='工作流执行失败时，先查看执行历史和节点错误日志，再检查模型连接。', **metadata):
    result = DocumentProcessor(collection).process_text(text, {'title': '排障指南', **metadata})
    assert result['success'], result
    return Document.objects.get(id=result['document_id'])


def test_fresh_collection_works_without_external_services(authenticated_client):
    response = authenticated_client.post(f'{BASE}/collections/', {'name': '无需部署'}, format='json')
    assert response.status_code == 201
    collection_id = response.data['id']
    response = authenticated_client.post(f'{BASE}/collections/{collection_id}/add_document/',
        {'title': '工作流排障', 'content': '工作流失败时请检查执行日志。', 'metadata': {'businessKey': 'ops'}}, format='json')
    assert response.status_code == 201
    document_id = response.data['document_id']
    response = authenticated_client.get(f'{BASE}/collections/{collection_id}/search/', {'query': '工作流失败'})
    assert response.status_code == 200
    assert response.data['results'][0]['metadata']['doc_id'] == document_id
    assert response.data['results'][0]['metadata']['businessKey'] == 'ops'
    assert response.data['results'][0]['citation'] == 1


def test_retrieval_survives_service_instances_and_does_not_leak_collections(collection, user):
    document = add(collection)
    other = DocumentCollection.objects.create(name='隔离库', vector_store=collection.vector_store, owner=user)
    add(other, '工作流秘密访问码 ONLY_OTHER_COLLECTION')
    results = RAGService(DocumentCollection.objects.get(pk=collection.id)).retrieve('工作流')
    assert {r['payload']['doc_id'] for r in results} == {str(document.id)}
    assert RAGService(collection)._get_collection_name() != RAGService(other)._get_collection_name()


def test_filters_threshold_and_no_match(collection):
    add(collection, category='ops')
    add(collection, '工作流设计涉及条件分支与节点配置。', category='design')
    assert len(RAGService(collection).retrieve('工作流', {'category': 'ops'})) == 1
    assert RAGService(collection).retrieve('xyzzymissingtoken') == []
    assert RAGService(collection, min_similarity=1).retrieve('工作流') == []
    with pytest.raises(ValueError):
        RAGService(collection).retrieve('test', {'category': {'$ne': 'ops'}})


@pytest.mark.parametrize('size,overlap', [(1, 0), (10, 9), (50, 20)])
@pytest.mark.parametrize('strategy', ['fixed', 'recursive'])
def test_chunking_preserves_all_characters_and_terminates(size, overlap, strategy):
    text = '段落一。段落二。\n\n文档末尾不应丢失。' * 4
    pieces = list(split_text(text, size, overlap, strategy))
    covered = set()
    for content, start, end in pieces:
        assert content == text[start:end]
        assert end - start <= size
        covered.update(range(start, end))
    assert all(index in covered for index, char in enumerate(text) if not char.isspace())
    assert pieces[-1][2] == len(text)


def test_invalid_chunk_overlap_rejected(collection):
    collection.chunk_overlap = collection.chunk_size
    with pytest.raises(ValueError):
        RAGService(collection).process_document('hello')


def test_reprocess_replaces_chunks_without_duplicating_document(collection, authenticated_client):
    document = add(collection)
    previous_ids = set(document.chunks.values_list('id', flat=True))
    for _ in range(2):
        response = authenticated_client.post(f'{BASE}/documents/{document.id}/reprocess/')
        assert response.status_code == 200
    assert collection.documents.count() == 1
    assert document.chunks.count() == 1
    assert not previous_ids & set(document.chunks.values_list('id', flat=True))
    collection.refresh_from_db()
    assert collection.document_count == collection.chunk_count == 1
    document.refresh_from_db()
    assert document.processed_at and document.status == 'completed'


def test_edit_reindexes_content_and_delete_removes_evidence(collection, authenticated_client):
    document = add(collection, 'olduniquevalue')
    response = authenticated_client.patch(f'{BASE}/documents/{document.id}/', {'content': 'newuniquevalue'}, format='json')
    assert response.status_code == 200
    assert not RAGService(collection).retrieve('olduniquevalue')
    assert RAGService(collection).retrieve('newuniquevalue')
    assert authenticated_client.delete(f'{BASE}/documents/{document.id}/').status_code == 204
    assert RAGService(collection).retrieve('newuniquevalue') == []
    collection.refresh_from_db()
    assert collection.chunk_count == collection.document_count == 0


def test_configuration_change_invalidates_old_index_and_reindex_recovers(collection, authenticated_client):
    add(collection)
    response = authenticated_client.patch(f'{BASE}/collections/{collection.id}/', {'chunk_size': 140}, format='json')
    assert response.status_code == 200
    assert not RAGService(collection).retrieve('工作流')
    response = authenticated_client.post(f'{BASE}/collections/{collection.id}/reindex/')
    assert response.data['completed'] == 1 and response.data['failed'] == 0
    collection.refresh_from_db()
    assert RAGService(collection).retrieve('工作流')


def test_task_uses_identical_processor_and_claim_prevents_duplicate_work(collection):
    from rag.tasks import process_document_async
    doc = Document.objects.create(collection=collection, title='异步处理', content='排查工作流执行日志')
    result = process_document_async(str(doc.id))
    assert result['success'] and result['chunks_created'] == 1
    doc.refresh_from_db()
    doc.status = 'processing'
    doc.save()
    assert not process_document_async(str(doc.id))['success']
    assert doc.chunks.count() == 1


def test_embedding_failure_is_not_fake_success(collection, authenticated_client):
    collection.retrieval_mode = 'semantic'
    collection.save()
    with patch.object(EmbeddingService, 'embed_texts', side_effect=RAGUnavailable('嵌入模型离线')):
        response = authenticated_client.post(f'{BASE}/collections/{collection.id}/add_document/',
            {'title': '离线测试', 'content': '工作流'}, format='json')
    assert response.status_code == 400
    document = collection.documents.get()
    assert document.status == 'failed' and '离线' in document.error_message
    assert not document.chunks.exists()


def test_embedding_without_credentials_never_returns_random_vectors(settings):
    settings.OPENAI_API_KEY = ''
    settings.RAG_EMBEDDING_API_KEY = ''
    with pytest.raises(RAGUnavailable, match='API Key'):
        EmbeddingService('openai:text-embedding-3-small').embed_text('工作流')


def test_semantic_vectors_persist_and_real_cosine_ranking_is_used(collection):
    collection.retrieval_mode = 'semantic'
    collection.save()
    with patch.object(EmbeddingService, 'embed_texts', side_effect=[[[1.0, 0.0]], [[0.0, 1.0]]]):
        relevant = add(collection, 'workflow logs')
        add(collection, 'unrelated cooking instructions')
    with patch.object(EmbeddingService, 'embed_text', return_value=[1.0, 0.0]):
        results = RAGService(DocumentCollection.objects.get(pk=collection.id)).retrieve('workflow')
    assert results[0]['payload']['doc_id'] == str(relevant.id)
    assert results[0]['score'] == 1
    assert relevant.chunks.get().embedding == [1.0, 0.0]


def test_hybrid_combines_both_rankings(collection):
    collection.retrieval_mode = 'hybrid'
    collection.save()
    with patch.object(EmbeddingService, 'embed_texts', side_effect=[[[1.0, 0.0]], [[0.7, 0.7]]]):
        relevant = add(collection, 'workflow execution logs')
        add(collection, 'language paraphrase')
    with patch.object(EmbeddingService, 'embed_text', return_value=[1.0, 0.0]):
        results = RAGService(collection).retrieve('workflow')
    assert results[0]['payload']['doc_id'] == str(relevant.id)
    assert results[0]['score'] > results[1]['score']


def test_empty_knowledge_never_invokes_llm(collection, authenticated_client):
    with patch('core.llm_factory.invoke_chat_messages') as llm:
        response = authenticated_client.post(f'{BASE}/collections/{collection.id}/ask/', {'query': '没有记录的问题'}, format='json')
    assert response.status_code == 200
    assert response.data['grounded'] is False
    assert response.data['answer_status'] == 'no_evidence'
    llm.assert_not_called()


def test_answer_contains_only_context_supplied_sources_and_counts_usage(collection, authenticated_client):
    document = add(collection)
    pipeline = RAGPipeline.objects.create(name='运维助手', collection=collection, owner=collection.owner)
    with patch('core.llm_factory.invoke_chat_messages', return_value='查看执行历史和节点日志。[1]') as llm:
        response = authenticated_client.post(f'{BASE}/pipelines/{pipeline.id}/query/', {'query': '工作流失败怎么办'}, format='json')
    assert response.status_code == 200
    assert response.data['sources'][0]['metadata']['doc_id'] == str(document.id)
    assert '[1]' in response.data['response']
    assert '不可信' in llm.call_args.args[0][0]['content']
    pipeline.refresh_from_db()
    assert pipeline.usage_count == 1


def test_generation_failure_is_http_503_not_a_fake_answer(collection, authenticated_client):
    add(collection)
    with patch('core.llm_factory.invoke_chat_messages', side_effect=RuntimeError('secret-provider-error')):
        response = authenticated_client.post(f'{BASE}/collections/{collection.id}/ask/', {'query': '工作流失败'}, format='json')
    assert response.status_code == 503
    assert 'secret-provider-error' not in str(response.data)


@pytest.mark.parametrize('params', [{'query': ''}, {'query': 'x', 'top_k': 'bad'}, {'query': 'x', 'top_k': 0},
                                  {'query': 'x', 'top_k': 999}, {'query': 'x', 'min_similarity': 2},
                                  {'query': 'x', 'filters': 'not json'}])
def test_search_validates_input(collection, authenticated_client, params):
    response = authenticated_client.get(f'{BASE}/collections/{collection.id}/search/', params)
    assert response.status_code == 400


def test_http_workflow_agent_all_enforce_current_access(collection, authenticated_client, user):
    from workflows.services.langchain_service import LangChainService
    from agents.engine import KnowledgeBaseTool
    stranger = get_user_model().objects.create_user(username='stranger', email='stranger@example.test')
    document = add(collection)
    other = DocumentCollection.objects.create(name='Private', vector_store=collection.vector_store, owner=stranger)
    assert authenticated_client.get(f'{BASE}/collections/{other.id}/').status_code == 404
    assert authenticated_client.post(f'{BASE}/pipelines/', {'name': 'unauthorized', 'collection': str(other.id)}, format='json').status_code == 400
    own_service = LangChainService(user_id=user.id)
    config = {'provider': 'knowledge_base', 'collectionId': str(collection.id), 'query': '工作流'}
    result = own_service.run_retriever_node(config, {})
    assert result['success'] and result['output']['sources'][0]['doc_id'] == str(document.id)
    assert '[1]' in result['output']['context']
    assert not LangChainService(user_id=stranger.id).run_retriever_node(config, {})['success']
    assert KnowledgeBaseTool(collection_ids=[str(collection.id)], user_id=user.id).run('工作流')['count'] == 1
    assert not KnowledgeBaseTool(collection_ids=[str(collection.id)], user_id=stranger.id).run('工作流')['success']


def test_api_key_is_write_only_and_empty_update_preserves_it(authenticated_client, collection):
    store = collection.vector_store
    response = authenticated_client.patch(f'{BASE}/vector-stores/{store.id}/', {'api_key': 'secret-value'}, format='json')
    assert response.status_code == 200 and 'secret-value' not in str(response.data)
    authenticated_client.patch(f'{BASE}/vector-stores/{store.id}/', {'api_key': ''}, format='json')
    store.refresh_from_db()
    assert store.api_key == 'secret-value'


@pytest.mark.parametrize('name,content', [('guide.md', '# 工作流\n检查日志'), ('guide.csv', '问题,处理\n失败,查日志'),
                                       ('guide.html', '<script>badsecret</script><p>检查日志</p>')])
def test_text_formats_parse_real_content(name, content):
    result = extract_file_content(SimpleUploadedFile(name, content.encode()))
    assert '日志' in result and 'badsecret' not in result


def test_docx_xlsx_pptx_extract_actual_content():
    from docx import Document as WordDocument
    from openpyxl import Workbook
    word = WordDocument()
    word.add_paragraph('工作流文档正文')
    buffer = BytesIO()
    word.save(buffer)
    assert '正文' in extract_file_content(SimpleUploadedFile('test.docx', buffer.getvalue()))
    workbook = Workbook()
    workbook.active.append(['故障', '查看日志'])
    buffer = BytesIO()
    workbook.save(buffer)
    assert '查看日志' in extract_file_content(SimpleUploadedFile('test.xlsx', buffer.getvalue()))
    buffer = BytesIO()
    with ZipFile(buffer, 'w') as archive:
        archive.writestr('ppt/slides/slide1.xml', '<root xmlns:a="urn:test"><a:t>幻灯片正文</a:t></root>')
    assert '幻灯片正文' in extract_file_content(SimpleUploadedFile('test.pptx', buffer.getvalue()))


@pytest.mark.parametrize('file', [SimpleUploadedFile('test.exe', b'binary'), SimpleUploadedFile('empty.txt', b''),
                                SimpleUploadedFile('test.txt', b'hello\x00world')])
def test_invalid_files_fail_explicitly(file):
    with pytest.raises(ValueError):
        extract_file_content(file)


@pytest.mark.parametrize('url', ['http://127.0.0.1/admin', 'http://169.254.169.254/latest/meta-data', 'http://[::1]/'])
def test_web_import_blocks_private_networks(url):
    with pytest.raises(ValueError, match='内网'):
        fetch_public_text(url)


def test_file_upload_marks_source_and_failed_files_create_no_documents(collection, authenticated_client):
    response = authenticated_client.post(f'{BASE}/collections/{collection.id}/upload_file/',
        {'file': SimpleUploadedFile('guide.md', '工作流运行指南'.encode())}, format='multipart')
    assert response.status_code == 201
    assert collection.documents.get().source_type == 'file'
    response = authenticated_client.post(f'{BASE}/collections/{collection.id}/upload_file/',
        {'file': SimpleUploadedFile('bad.exe', b'not text')}, format='multipart')
    assert response.status_code == 400 and collection.documents.count() == 1


def test_system_guides_import_is_idempotent(collection, authenticated_client):
    response = authenticated_client.post(f'{BASE}/collections/{collection.id}/import_guides/')
    assert response.status_code == 200 and response.data['failed'] == 0
    ids = set(collection.documents.values_list('id', flat=True))
    response = authenticated_client.post(f'{BASE}/collections/{collection.id}/import_guides/')
    assert response.data['failed'] == 0
    assert set(collection.documents.values_list('id', flat=True)) == ids
    assert all(r['unchanged'] for r in response.data['results'])


def test_qdrant_preserves_previous_documents_and_deletes_only_selected_vectors():
    from qdrant_client import QdrantClient
    from rag.services import QdrantAdapter
    import uuid
    client = QdrantClient(':memory:')
    adapter = QdrantAdapter(client, None)
    first, second = str(uuid.uuid4()), str(uuid.uuid4())
    try:
        assert adapter.connect()
        adapter.create_collection('flowly_test', 2)
        adapter.upsert('flowly_test', [{'id': first, 'vector': [1.0, 0.0], 'payload': {'content': 'first'}}])
        adapter.create_collection('flowly_test', 2)
        adapter.upsert('flowly_test', [{'id': second, 'vector': [0.0, 1.0], 'payload': {'content': 'second'}}])
        assert {r['id'] for r in adapter.search('flowly_test', [1.0, 0.0], 5)} == {first, second}
        adapter.delete('flowly_test', [first])
        assert [r['id'] for r in adapter.search('flowly_test', [1.0, 0.0], 5)] == [second]
    finally:
        adapter.disconnect()


def test_qdrant_local_storage_survives_client_restarts_and_reprocessing(collection, tmp_path):
    store = collection.vector_store
    store.store_type = 'qdrant'
    store.path = str(tmp_path / 'qdrant')
    store.save()
    collection.retrieval_mode = 'semantic'
    collection.save()
    with patch.object(EmbeddingService, 'embed_texts', return_value=[[1.0, 0.0]]):
        document = add(collection, '持久化检索')
        assert DocumentProcessor(collection).process_document(document)['success']
    with patch.object(EmbeddingService, 'embed_text', return_value=[1.0, 0.0]):
        assert len(RAGService(collection).retrieve('检索')) == 1
        assert len(RAGService(DocumentCollection.objects.get(id=collection.id)).retrieve('检索')) == 1
    document.delete()
    assert not collection.documents.exists()


def test_chroma_persistent_adapter_uses_similarity_not_distance(tmp_path):
    import chromadb
    import uuid
    from rag.services import ChromaAdapter
    client = chromadb.PersistentClient(path=str(tmp_path / 'chroma'))
    adapter = ChromaAdapter(client, None)
    key = str(uuid.uuid4())
    adapter.create_collection('flowly_test', 2)
    adapter.upsert('flowly_test', [{'id': key, 'vector': [1.0, 0.0], 'payload': {'content': 'source', 'nested': {'tag': 'ops'}}}])
    result = adapter.search('flowly_test', [1.0, 0.0])[0]
    assert result['score'] == pytest.approx(1.0)
    assert result['payload']['nested']['tag'] == 'ops'
    adapter.delete('flowly_test', [key])
    assert not adapter.search('flowly_test', [1.0, 0.0])


def test_membership_revocation_blocks_previously_bound_pipeline_and_agent(collection, user):
    from tenants.models import Tenant, TenantMembership
    from agents.engine import KnowledgeBaseTool
    from rest_framework.test import APIClient
    member = get_user_model().objects.create_user(username='team-member', email='member@example.test')
    tenant = Tenant.objects.create(name='Team', slug='rag-team', owner=user)
    membership, _ = TenantMembership.objects.get_or_create(tenant=tenant, user=member, defaults={'role': 'member'})
    collection.tenant = tenant
    collection.save()
    document = add(collection)
    pipeline = RAGPipeline.objects.create(name='shared', collection=collection, owner=member)
    client = APIClient()
    client.force_authenticate(member)
    assert client.get(f'{BASE}/pipelines/{pipeline.id}/').status_code == 200
    assert client.delete(f'{BASE}/documents/{document.id}/').status_code == 403
    membership.is_active = False
    membership.save()
    assert client.get(f'{BASE}/pipelines/{pipeline.id}/').status_code == 404
    assert not KnowledgeBaseTool(collection_ids=[str(collection.id)], user_id=member.id).run('工作流')['success']


def test_embedding_http_contract_preserves_order_and_raises_on_outage(settings):
    import httpx
    settings.RAG_EMBEDDING_BASE_URL = 'https://embedding.example/v1'
    settings.RAG_EMBEDDING_API_KEY = 'test-only'
    def handler(request):
        assert request.url.path == '/v1/embeddings'
        return httpx.Response(200, json={'data': [{'index': 1, 'embedding': [0.0, 1.0]}, {'index': 0, 'embedding': [1.0, 0.0]}]})
    client = httpx.Client(transport=httpx.MockTransport(handler))
    with patch('rag.pipeline.httpx.Client', return_value=client):
        assert EmbeddingService('openai:test').embed_texts(['first', 'second']) == [[1.0, 0.0], [0.0, 1.0]]


def test_url_import_preserves_provenance(collection, authenticated_client):
    with patch('rag.pipeline.fetch_public_text', return_value=('公开页面中的工作流说明', 'https://example.com/final')):
        response = authenticated_client.post(f'{BASE}/collections/{collection.id}/add_url/',
            {'url': 'https://example.com/source', 'title': '网页指南'}, format='json')
    assert response.status_code == 201
    document = collection.documents.get()
    assert document.source_type == 'url' and document.source_url == 'https://example.com/final'
    assert RAGService(collection).retrieve('工作流')[0]['payload']['source_url'] == document.source_url


def test_crashed_processing_can_be_retried_after_lease_expiry(collection, settings):
    from datetime import timedelta
    from django.utils import timezone
    settings.RAG_PROCESSING_LEASE_SECONDS = 60
    document = add(collection)
    Document.objects.filter(pk=document.pk).update(status='processing', updated_at=timezone.now() - timedelta(seconds=61))
    assert DocumentProcessor(collection).process_document(document)['success']
    assert collection.documents.count() == 1
    assert len(RAGService(collection).retrieve('工作流')) == 1


def test_stale_processor_cannot_overwrite_newer_claim(collection):
    from datetime import timedelta
    from django.utils import timezone
    document = add(collection)
    processor = DocumentProcessor(collection)
    original = processor.rag_service.process_document
    def superseded(*args, **kwargs):
        vectors = original(*args, **kwargs)
        Document.objects.filter(pk=document.pk).update(updated_at=timezone.now() + timedelta(seconds=1), error_message='new worker')
        return vectors
    with patch.object(processor.rag_service, 'process_document', side_effect=superseded):
        assert not processor.process_document(document)['success']
    document.refresh_from_db()
    assert document.status == 'processing' and document.error_message == 'new worker'


def test_changing_embedding_dimensions_rebuilds_external_index(collection, tmp_path):
    from rag.services import get_vector_store_client, PineconeAdapter
    store = collection.vector_store
    store.store_type, store.path = 'qdrant', str(tmp_path / 'model-change')
    store.save()
    collection.retrieval_mode, collection.embedding_model = 'semantic', 'openai:old'
    collection.save()
    with patch.object(EmbeddingService, 'embed_texts', return_value=[[1.0, 0.0]]):
        document = add(collection)
    old_namespace = RAGService(collection)._get_collection_name()
    collection.embedding_model = 'openai:new'
    collection.save()
    with patch.object(EmbeddingService, 'embed_texts', return_value=[[1.0, 0.0, 0.0]]):
        assert DocumentProcessor(collection).process_document(document)['success']
    new_namespace = RAGService(collection)._get_collection_name()
    assert old_namespace != new_namespace
    assert len(PineconeAdapter.name(new_namespace)) <= 45
    with patch.object(EmbeddingService, 'embed_text', return_value=[1.0, 0.0, 0.0]):
        assert len(RAGService(collection).retrieve('工作流')) == 1
    client = get_vector_store_client(store)
    try:
        assert not client.search(old_namespace, [1.0, 0.0])
    finally:
        client.disconnect()
    document.delete()
