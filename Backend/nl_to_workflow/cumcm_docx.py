# -*- coding: utf-8 -*-
"""
Convert CUMCM paper Markdown to Word (.docx).

Used by file_writer (format=docx) and CUMCM pipeline export node.
Follows prompts/docx_write_markdown.md heading and keyword conventions.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import List, Optional, Union


def _strip_md_bold(text: str) -> tuple:
    """Return (label, rest) for **label**：content patterns."""
    m = re.match(r'^\*\*(.+?)\*\*[：:]\s*(.*)$', text.strip(), re.S)
    if m:
        return m.group(1), m.group(2)
    return '', text


def _add_runs(paragraph, text: str) -> None:
    """Add text with simple **bold** segments."""
    parts = re.split(r'(\*\*.+?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part:
            paragraph.add_run(part)


def markdown_to_docx(
    markdown_text: str,
    output_path: Union[str, Path],
    *,
    title: Optional[str] = None,
) -> Path:
    """
    Convert Markdown paper to .docx (headings, paragraphs, simple tables, block math as text).

    Args:
        markdown_text: Full paper Markdown
        output_path: Destination .docx path
        title: Optional override for document title (else first # line)

    Returns:
        Path to written file
    """
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(12)

    lines = (markdown_text or '').splitlines()
    i = 0
    doc_title = title

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped.startswith('> '):
            p = doc.add_paragraph(stripped[2:])
            p.runs[0].italic = True
            i += 1
            continue

        if stripped.startswith('# ') and not stripped.startswith('## '):
            doc_title = doc_title or stripped[2:].strip()
            h = doc.add_heading(doc_title, level=0)
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        if stripped.startswith('## '):
            doc.add_heading(stripped[3:].strip(), level=1)
            i += 1
            continue

        if stripped.startswith('### '):
            doc.add_heading(stripped[4:].strip(), level=2)
            i += 1
            continue

        if stripped.startswith('#### '):
            doc.add_heading(stripped[5:].strip(), level=3)
            i += 1
            continue

        if stripped.startswith('$$'):
            block = []
            if stripped.endswith('$$') and len(stripped) > 4:
                block.append(stripped[2:-2].strip())
                i += 1
            else:
                i += 1
                while i < len(lines) and not lines[i].strip().endswith('$$'):
                    block.append(lines[i])
                    i += 1
                if i < len(lines):
                    block.append(lines[i].strip().rstrip('$'))
                    i += 1
            p = doc.add_paragraph()
            run = p.add_run(' '.join(block))
            run.italic = True
            continue

        if stripped.startswith('|') and '|' in stripped[1:]:
            table_lines: List[str] = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1
            if len(table_lines) >= 2:
                header = [c.strip() for c in table_lines[0].strip('|').split('|')]
                rows = []
                for tl in table_lines[2:]:
                    if re.match(r'^\|?\s*[-:| ]+\|', tl):
                        continue
                    rows.append([c.strip() for c in tl.strip('|').split('|')])
                if header:
                    tbl = doc.add_table(rows=1 + len(rows), cols=len(header))
                    tbl.style = 'Table Grid'
                    for ci, h in enumerate(header):
                        tbl.rows[0].cells[ci].text = h
                    for ri, row in enumerate(rows):
                        for ci, cell in enumerate(row[: len(header)]):
                            tbl.rows[ri + 1].cells[ci].text = cell
            continue

        if stripped.startswith('!['):
            alt_m = re.match(r'!\[(.*?)\]\((.*?)\)', stripped)
            if alt_m:
                p = doc.add_paragraph()
                p.add_run(f'[{alt_m.group(1)}]').italic = True
                p.add_run(f' ({alt_m.group(2)})')
            i += 1
            continue

        label, body = _strip_md_bold(stripped)
        p = doc.add_paragraph()
        if label:
            run = p.add_run(label)
            run.bold = True
            p.add_run('：' + body)
        else:
            _add_runs(p, stripped)
        i += 1

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    return out
