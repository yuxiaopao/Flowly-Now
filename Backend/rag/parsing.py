"""Bounded document extraction and public-web ingestion."""
import ipaddress
import socket
from io import BytesIO
from pathlib import Path
from urllib.parse import urlsplit, urljoin
from zipfile import ZipFile, is_zipfile

import httpx
from bs4 import BeautifulSoup
from django.conf import settings
from defusedxml import ElementTree

from .constants import (
    FILE_EXTENSIONS, TEXT_ENCODINGS, DEFAULT_MAX_FILE_BYTES, DEFAULT_MAX_TEXT_CHARS,
    DEFAULT_REQUEST_TIMEOUT,
)


def validate_content(content):
    if not isinstance(content, str) or not content.strip():
        raise ValueError('文档没有可提取的文本；扫描件需先进行 OCR。')
    if len(content) > getattr(settings, 'RAG_MAX_TEXT_CHARS', DEFAULT_MAX_TEXT_CHARS):
        raise ValueError('文档文本过长，请拆分后上传。')
    if '\x00' in content:
        raise ValueError('文件包含二进制内容，请上传受支持的文档格式。')
    return content.strip()


def decode_text(raw):
    encodings = TEXT_ENCODINGS if raw.startswith((b'\xff\xfe', b'\xfe\xff')) else (
        'utf-8-sig', 'gb18030',
    )
    for encoding in encodings:
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError('无法识别文本编码，请转为 UTF-8 后上传。')


def html_text(content):
    soup = BeautifulSoup(content, 'html.parser')
    for element in soup(['script', 'style', 'noscript', 'nav', 'footer']):
        element.decompose()
    return soup.get_text(separator='\n', strip=True)


def extract_file_content(file):
    suffix = Path(file.name).suffix.lower()
    if suffix not in FILE_EXTENSIONS:
        raise ValueError('不支持此格式。支持 TXT、Markdown、CSV、JSON、HTML、PDF、DOCX、XLSX、PPTX。')
    limit = getattr(settings, 'RAG_MAX_FILE_BYTES', DEFAULT_MAX_FILE_BYTES)
    if getattr(file, 'size', 0) > limit:
        raise ValueError('文件超过上传大小限制，请拆分后上传。')
    raw = file.read(limit + 1)
    if len(raw) > limit:
        raise ValueError('文件超过上传大小限制，请拆分后上传。')
    stream = BytesIO(raw)
    if suffix in {'.docx', '.xlsx', '.pptx'}:
        if not is_zipfile(stream):
            raise ValueError('Office 文件损坏或格式与扩展名不符。')
        with ZipFile(stream) as archive:
            if sum(entry.file_size for entry in archive.infolist()) > limit * 5:
                raise ValueError('Office 文档解压后过大，请拆分后上传。')
        stream.seek(0)
    if suffix == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(stream)
        if reader.is_encrypted:
            raise ValueError('请先解除 PDF 密码保护。')
        content = '\n\n'.join(page.extract_text() or '' for page in reader.pages)
    elif suffix == '.docx':
        from docx import Document
        document = Document(stream)
        parts = [paragraph.text for paragraph in document.paragraphs]
        parts.extend('\t'.join(cell.text for cell in row.cells)
                     for table in document.tables for row in table.rows)
        content = '\n'.join(parts)
    elif suffix == '.xlsx':
        from openpyxl import load_workbook
        workbook = load_workbook(stream, read_only=True, data_only=True)
        try:
            parts, length = [], 0
            for sheet in workbook:
                parts.append(f'[{sheet.title}]')
                for row in sheet.iter_rows(values_only=True):
                    line = '\t'.join('' if cell is None else str(cell) for cell in row)
                    length += len(line) + 1
                    if length > getattr(settings, 'RAG_MAX_TEXT_CHARS', DEFAULT_MAX_TEXT_CHARS):
                        raise ValueError('表格文本过长，请拆分后上传。')
                    parts.append(line)
            content = '\n'.join(parts)
        finally:
            workbook.close()
    elif suffix == '.pptx':
        with ZipFile(stream) as archive:
            import re
            names = sorted((name for name in archive.namelist()
                            if re.fullmatch(r'ppt/slides/slide\d+\.xml', name)),
                           key=lambda name: int(re.search(r'(\d+)\.xml', name)[1]))
            content = '\n\n'.join('\n'.join(
                node.text or '' for node in ElementTree.fromstring(archive.read(name)).iter()
                if node.tag.endswith('}t')) for name in names)
    else:
        content = decode_text(raw)
        if suffix in {'.html', '.htm'}:
            content = html_text(content)
    return validate_content(content)


def fetch_public_text(url):
    """Validate and pin every redirect target, avoiding DNS rebinding and proxy bypass."""
    timeout = getattr(settings, 'RAG_REQUEST_TIMEOUT', DEFAULT_REQUEST_TIMEOUT)
    limit = getattr(settings, 'RAG_MAX_FILE_BYTES', DEFAULT_MAX_FILE_BYTES)
    with httpx.Client(timeout=timeout, trust_env=False, follow_redirects=False) as client:
        for _ in range(6):
            parsed = urlsplit(url)
            if parsed.scheme not in {'http', 'https'} or not parsed.hostname or parsed.username or parsed.password:
                raise ValueError('仅支持不含登录凭证的 HTTP/HTTPS 网页。')
            host = parsed.hostname
            addresses = socket.getaddrinfo(host, parsed.port or (443 if parsed.scheme == 'https' else 80), type=socket.SOCK_STREAM)
            if not addresses or any(not ipaddress.ip_address(a[4][0]).is_global for a in addresses):
                raise ValueError('不允许抓取内网、回环或保留网络地址。')
            request = client.build_request('GET', url)
            request.url = request.url.copy_with(host=addresses[0][4][0])
            request.headers['Host'] = parsed.netloc
            request.extensions['sni_hostname'] = host
            response = client.send(request, stream=True)
            try:
                if response.is_redirect:
                    location = response.headers.get('location')
                    if not location:
                        raise ValueError('网页重定向缺少目标地址。')
                    url = urljoin(url, location)
                    continue
                response.raise_for_status()
                content_type = response.headers.get('content-type', '').lower()
                if not any(kind in content_type for kind in ('text/', 'application/json', 'application/xhtml+xml')):
                    raise ValueError('URL 必须指向文本或 HTML 网页；其他文档请下载后上传。')
                chunks, size = [], 0
                for chunk in response.iter_bytes():
                    size += len(chunk)
                    if size > limit:
                        raise ValueError('网页内容超过大小限制。')
                    chunks.append(chunk)
                content = decode_text(b''.join(chunks))
                return validate_content(html_text(content) if 'html' in content_type else content), url
            finally:
                response.close()
    raise ValueError('网页重定向次数过多。')
